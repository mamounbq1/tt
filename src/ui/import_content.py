"""
Import Content Frame
Import courses and homework from Excel/CSV files
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import logging
import csv
import os
from src.utils.theme import ThemeManager
from src.utils.config import DB_PATH


class ImportContentFrame(ttk.Frame):
    """Frame for importing content from Excel/CSV files"""
    
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        self.db = controller.database
        
        self.imported_data = []
        self.file_path = None
        
        self.build_ui()
    
    def build_ui(self):
        """Build the user interface"""
        # Main container
        main_container = ttk.Frame(self, padding=10)
        main_container.pack(fill='both', expand=True)
        
        # Top bar
        top_bar = ttk.Frame(main_container)
        top_bar.pack(fill='x', pady=(0, 10))
        
        back_btn = ThemeManager.create_button(
            top_bar,
            text='← Retour',
            command=self.go_back
        )
        back_btn.pack(side='left')
        
        title = ThemeManager.create_label(
            top_bar,
            text='📥 Importer Contenu',
            style='Heading.TLabel'
        )
        title.pack(side='left', padx=20)
        
        # Instructions
        instructions = ttk.LabelFrame(main_container, text='Instructions', padding=10)
        instructions.pack(fill='x', pady=(0, 10))
        
        inst_text = """
Format du fichier CSV/Excel:
- Colonne 1: Date (JJ/MM/AAAA)
- Colonne 2: Classe
- Colonne 3: Matière
- Colonne 4: Contenu du cours
- Colonne 5: Devoirs (optionnel)
- Colonne 6: Examen (optionnel)

Exemple:
15/01/2026,6ème A,Mathématiques,Géométrie - Les triangles,Exercices page 45,
20/01/2026,5ème B,Français,Grammaire - Les pronoms,Rédaction,Contrôle vendredi
        """
        
        ttk.Label(
            instructions,
            text=inst_text.strip(),
            font=('Arial', 9),
            justify='left'
        ).pack(anchor='w')
        
        # File selection
        file_frame = ttk.Frame(main_container)
        file_frame.pack(fill='x', pady=(0, 10))
        
        select_btn = ThemeManager.create_button(
            file_frame,
            text='📁 Sélectionner Fichier',
            command=self.select_file
        )
        select_btn.pack(side='left', padx=5)
        
        self.file_label = ttk.Label(
            file_frame,
            text='Aucun fichier sélectionné',
            font=('Arial', 9, 'italic')
        )
        self.file_label.pack(side='left', padx=10)
        
        # Action buttons
        action_bar = ttk.Frame(main_container)
        action_bar.pack(fill='x', pady=(0, 10))
        
        preview_btn = ThemeManager.create_button(
            action_bar,
            text='👁️ Aperçu',
            command=self.preview_data
        )
        preview_btn.pack(side='left', padx=5)
        
        import_btn = ThemeManager.create_button(
            action_bar,
            text='✓ Importer',
            command=self.import_data
        )
        import_btn.pack(side='left', padx=5)
        
        clear_btn = ThemeManager.create_button(
            action_bar,
            text='🗑️ Effacer',
            command=self.clear_preview
        )
        clear_btn.pack(side='left', padx=5)
        
        # Preview area
        preview_label = ttk.Label(main_container, text='Aperçu des données:', font=('Arial', 10, 'bold'))
        preview_label.pack(anchor='w', pady=(0, 5))
        
        # Treeview for preview
        columns = ('date', 'classe', 'matiere', 'contenu', 'devoirs', 'examen')
        self.preview_tree = ttk.Treeview(main_container, columns=columns, show='headings', height=15)
        
        self.preview_tree.heading('date', text='Date')
        self.preview_tree.heading('classe', text='Classe')
        self.preview_tree.heading('matiere', text='Matière')
        self.preview_tree.heading('contenu', text='Contenu')
        self.preview_tree.heading('devoirs', text='Devoirs')
        self.preview_tree.heading('examen', text='Examen')
        
        self.preview_tree.column('date', width=100)
        self.preview_tree.column('classe', width=100)
        self.preview_tree.column('matiere', width=120)
        self.preview_tree.column('contenu', width=200)
        self.preview_tree.column('devoirs', width=150)
        self.preview_tree.column('examen', width=100)
        
        # Scrollbars
        v_scroll = ttk.Scrollbar(main_container, orient='vertical', command=self.preview_tree.yview)
        h_scroll = ttk.Scrollbar(main_container, orient='horizontal', command=self.preview_tree.xview)
        self.preview_tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)
        
        self.preview_tree.pack(side='left', fill='both', expand=True)
        v_scroll.pack(side='right', fill='y')
        h_scroll.pack(side='bottom', fill='x')
        
        # Status bar
        self.status_label = ttk.Label(
            self,
            text='Prêt à importer',
            font=('Arial', 9, 'italic'),
            foreground='green'
        )
        self.status_label.pack(side='bottom', anchor='w', padx=10, pady=5)
    
    def select_file(self):
        """Select CSV/Excel file to import"""
        filetypes = [
            ('CSV files', '*.csv'),
            ('Excel files', '*.xlsx *.xls'),
            ('All files', '*.*')
        ]
        
        filename = filedialog.askopenfilename(
            title='Sélectionner un fichier',
            filetypes=filetypes
        )
        
        if filename:
            self.file_path = filename
            self.file_label.config(text=os.path.basename(filename))
            self.status_label.config(text=f'Fichier sélectionné: {os.path.basename(filename)}', foreground='blue')
            logging.info(f"Selected file: {filename}")
    
    def preview_data(self):
        """Preview data from selected file"""
        if not self.file_path:
            messagebox.showwarning("Attention", "Veuillez d'abord sélectionner un fichier")
            return
        
        try:
            # Clear existing preview
            for item in self.preview_tree.get_children():
                self.preview_tree.delete(item)
            
            self.imported_data = []
            
            # Read CSV file
            if self.file_path.endswith('.csv'):
                with open(self.file_path, 'r', encoding='utf-8') as f:
                    reader = csv.reader(f)
                    
                    # Skip header if present
                    first_row = next(reader, None)
                    if first_row and not first_row[0].strip().replace('/', '').isdigit():
                        # Has header, skip it
                        pass
                    else:
                        # No header, process first row
                        if first_row:
                            self.imported_data.append(first_row)
                    
                    # Read remaining rows
                    for row in reader:
                        if len(row) >= 4:  # At least date, classe, matiere, contenu
                            self.imported_data.append(row)
            
            elif self.file_path.endswith(('.xlsx', '.xls')):
                # Try to import openpyxl for Excel
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(self.file_path)
                    ws = wb.active
                    
                    # Skip header row if present
                    start_row = 2 if ws.cell(1, 1).value and not str(ws.cell(1, 1).value).replace('/', '').isdigit() else 1
                    
                    for row in ws.iter_rows(min_row=start_row, values_only=True):
                        if row[0]:  # Has date
                            self.imported_data.append([str(cell) if cell else '' for cell in row])
                    
                except ImportError:
                    messagebox.showerror(
                        "Erreur",
                        "Le module 'openpyxl' n'est pas installé.\n"
                        "Veuillez utiliser un fichier CSV ou installer openpyxl:\n"
                        "pip install openpyxl"
                    )
                    return
            
            # Display in treeview
            for row in self.imported_data:
                # Ensure row has 6 columns
                while len(row) < 6:
                    row.append('')
                
                self.preview_tree.insert('', 'end', values=row[:6])
            
            count = len(self.imported_data)
            self.status_label.config(
                text=f'Aperçu: {count} entrée(s) trouvée(s)',
                foreground='green'
            )
            logging.info(f"Previewed {count} entries from {self.file_path}")
            
        except Exception as e:
            logging.error(f"Error previewing file: {e}")
            messagebox.showerror("Erreur", f"Impossible de lire le fichier:\n{str(e)}")
            self.status_label.config(text='Erreur lors de la lecture du fichier', foreground='red')
    
    def import_data(self):
        """Import data into database"""
        if not self.imported_data:
            messagebox.showwarning("Attention", "Aucune donnée à importer. Veuillez d'abord prévisualiser.")
            return
        
        # Confirm import
        result = messagebox.askyesno(
            "Confirmation",
            f"Voulez-vous importer {len(self.imported_data)} entrée(s) dans la base de données?"
        )
        
        if not result:
            return
        
        try:
            from datetime import datetime
            
            # Get default teacher ID
            teachers = self.db.execute_query("SELECT id FROM teachers LIMIT 1")
            teacher_id = teachers[0]['id'] if teachers else 1
            
            imported_count = 0
            errors = []
            
            for idx, row in enumerate(self.imported_data, start=1):
                try:
                    # Parse data
                    date_str = row[0].strip()
                    classe = row[1].strip()
                    matiere = row[2].strip()
                    contenu = row[3].strip()
                    devoirs = row[4].strip() if len(row) > 4 else ''
                    examen = row[5].strip() if len(row) > 5 else ''
                    
                    # Validate required fields
                    if not date_str or not classe or not matiere or not contenu:
                        errors.append(f"Ligne {idx}: Champs requis manquants")
                        continue
                    
                    # Convert date format
                    try:
                        date_obj = datetime.strptime(date_str, '%d/%m/%Y')
                        formatted_date = date_obj.strftime('%Y-%m-%d')
                    except ValueError:
                        errors.append(f"Ligne {idx}: Format de date invalide: {date_str}")
                        continue
                    
                    # Insert into database
                    query = """
                        INSERT INTO homework_entries 
                        (date, class_name, subject, content, homework, exam, teacher_id)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                    """
                    self.db.execute_update(
                        query,
                        (formatted_date, classe, matiere, contenu, devoirs, examen, teacher_id)
                    )
                    
                    imported_count += 1
                    
                except Exception as e:
                    errors.append(f"Ligne {idx}: {str(e)}")
            
            # Show results
            if errors:
                error_msg = f"{imported_count}/{len(self.imported_data)} entrées importées.\n\nErreurs:\n"
                error_msg += '\n'.join(errors[:10])  # Show first 10 errors
                if len(errors) > 10:
                    error_msg += f"\n... et {len(errors) - 10} autres erreurs"
                messagebox.showwarning("Import Partiel", error_msg)
            else:
                messagebox.showinfo(
                    "Succès",
                    f"{imported_count} entrée(s) importée(s) avec succès!"
                )
            
            self.status_label.config(
                text=f'Import terminé: {imported_count}/{len(self.imported_data)} entrées importées',
                foreground='green' if not errors else 'orange'
            )
            
            logging.info(f"Imported {imported_count} entries from {self.file_path}")
            
        except Exception as e:
            logging.error(f"Error importing data: {e}")
            messagebox.showerror("Erreur", f"Erreur lors de l'importation:\n{str(e)}")
            self.status_label.config(text='Erreur lors de l\'importation', foreground='red')
    
    def clear_preview(self):
        """Clear preview data"""
        for item in self.preview_tree.get_children():
            self.preview_tree.delete(item)
        
        self.imported_data = []
        self.file_path = None
        self.file_label.config(text='Aucun fichier sélectionné')
        self.status_label.config(text='Aperçu effacé', foreground='blue')
        logging.info("Cleared import preview")
    
    def go_back(self):
        """Return to dashboard"""
        self.controller.show_frame('DashboardFrame')
